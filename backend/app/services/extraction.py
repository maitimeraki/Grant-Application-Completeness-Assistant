from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass, field


@dataclass
class BlockData:
    para_index: int
    text: str
    page: int | None = None
    section_path: str | None = None


@dataclass
class ExtractionResult:
    blocks: list[BlockData]
    warnings: list[str] = field(default_factory=list)


HEADING_PATTERNS = [
    re.compile(r"^#{1,6}\s+(?P<t>.+)$"),
    re.compile(r"^(?P<t>\d+(\.\d+){0,3})\s+(?P<title>[A-Z].{2,80})$"),
    re.compile(r"^(?P<t>[A-Z][A-Z0-9 ,&/()\-\']{3,79})$"),
]
SKIP_LINES = re.compile(r"^page \d+ of \d+$", re.IGNORECASE)


def content_hash(data: bytes, *, is_text: bool = False) -> str:
    if is_text:
        data = b"text:" + data
    return hashlib.sha256(data).hexdigest()


def _split_paragraphs(text: str) -> list[str]:
    parts = re.split(r"\n\s*\n", text)
    out = []
    for part in parts:
        cleaned = "\n".join(
            line for line in part.splitlines() if not SKIP_LINES.match(line.strip())
        ).strip()
        if cleaned:
            out.append(cleaned)
    return out


def _heading_of(line: str) -> str | None:
    stripped = line.strip()
    if not stripped or len(stripped) > 120:
        return None
    for pattern in HEADING_PATTERNS:
        m = pattern.match(stripped)
        if m:
            title = m.groupdict().get("title") or m.groupdict().get("t") or stripped
            return title.strip()
    return None


def _build_blocks(paragraphs: list[tuple[str, int | None]]) -> list[BlockData]:
    blocks: list[BlockData] = []
    section_stack: list[str] = []
    idx = 0
    for para, page in paragraphs:
        first_line = para.splitlines()[0]
        heading = _heading_of(first_line)
        if heading and len(para.splitlines()) <= 3:
            level = 1 if not section_stack else 2
            section_stack = section_stack[: level - 1] + [heading]
            section_path = " > ".join(section_stack)
        else:
            section_path = " > ".join(section_stack) if section_stack else None
        blocks.append(BlockData(para_index=idx, text=para, page=page, section_path=section_path))
        idx += 1
    return blocks


def extract_from_pdf(data: bytes) -> ExtractionResult:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    paragraphs: list[tuple[str, int | None]] = []
    for page_no, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        for para in _split_paragraphs(text):
            paragraphs.append((para, page_no))

    warnings: list[str] = []
    total_chars = sum(len(p) for p, _ in paragraphs)
    if not paragraphs:
        warnings.append(
            "No text could be extracted from this PDF. Scanned/image-only PDFs "
            "are not supported (OCR is out of scope) — please supply a "
            "text-bearing file or paste the text."
        )
    elif reader.pages and total_chars / len(reader.pages) < 60:
        warnings.append(
            "Very little text was extracted per page — this may be a scanned "
            "PDF. Results may be incomplete (OCR is out of scope)."
        )
    return ExtractionResult(blocks=_build_blocks(paragraphs), warnings=warnings)


def extract_from_docx(data: bytes) -> ExtractionResult:
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(data))
    paragraphs: list[tuple[str, int | None]] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append((text, None))
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paragraphs.append((" | ".join(cells), None))

    warnings: list[str] = []
    if not paragraphs:
        warnings.append("No text could be extracted from this DOCX file.")
    return ExtractionResult(blocks=_build_blocks(paragraphs), warnings=warnings)


def extract_from_text(data: bytes) -> ExtractionResult:
    text = data.decode("utf-8", errors="replace")
    paragraphs = [(para, None) for para in _split_paragraphs(text)]
    warnings: list[str] = []
    if not paragraphs:
        warnings.append("The pasted/uploaded text is empty.")
    return ExtractionResult(blocks=_build_blocks(paragraphs), warnings=warnings)


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".txt", ".markdown"}


def extract_from_file(filename: str, data: bytes) -> ExtractionResult:
    lower = (filename or "").lower()
    if lower.endswith(".pdf"):
        return extract_from_pdf(data)
    if lower.endswith(".docx"):
        return extract_from_docx(data)
    if lower.endswith((".md", ".txt", ".markdown")):
        return extract_from_text(data)
    raise ValueError(
        f"Unsupported file type for '{filename}'. Supported: PDF, DOCX, MD, TXT."
    )
